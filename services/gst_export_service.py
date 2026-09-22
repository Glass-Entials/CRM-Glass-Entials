"""
GST Export Service - GlassEntials CRM
Reads data exclusively from the existing Invoice module.
"""
from collections import defaultdict
from datetime import date, datetime
from typing import Dict, List, Optional

from sqlalchemy.orm import joinedload
from model import db, Invoice, InvoiceItem, InvoiceStatus, Customer, Lead, QuotationSettings


GSTIN_STATE_CODES = {
    "01": "Jammu & Kashmir", "02": "Himachal Pradesh", "03": "Punjab",
    "04": "Chandigarh", "05": "Uttarakhand", "06": "Haryana",
    "07": "Delhi", "08": "Rajasthan", "09": "Uttar Pradesh",
    "10": "Bihar", "11": "Sikkim", "12": "Arunachal Pradesh",
    "13": "Nagaland", "14": "Manipur", "15": "Mizoram",
    "16": "Tripura", "17": "Meghalaya", "18": "Assam",
    "19": "West Bengal", "20": "Jharkhand", "21": "Odisha",
    "22": "Chhattisgarh", "23": "Madhya Pradesh", "24": "Gujarat",
    "25": "Daman & Diu", "26": "Dadra & Nagar Haveli", "27": "Maharashtra",
    "28": "Andhra Pradesh", "29": "Karnataka", "30": "Goa",
    "31": "Lakshadweep", "32": "Kerala", "33": "Tamil Nadu",
    "34": "Puducherry", "35": "Andaman & Nicobar Islands",
    "36": "Telangana", "37": "Andhra Pradesh (New)",
    "38": "Ladakh", "97": "Other Territory", "99": "Centre Jurisdiction",
}

STATE_NAME_TO_CODE = {v: k for k, v in GSTIN_STATE_CODES.items()}

UQC_MAP = {
    "sq.ft": "SQF", "sqft": "SQF", "sqm": "SQM", "m": "MTR",
    "ft": "FT", "pcs": "PCS", "piece": "PCS", "nos": "NOS",
    "kg": "KGS", "ltr": "LTR", "box": "BOX",
}


def _uqc(unit):
    if not unit:
        return "OTH"
    return UQC_MAP.get(unit.lower().strip(), "OTH")


def _state_from_gstin(gstin):
    if gstin and len(gstin) >= 2:
        return GSTIN_STATE_CODES.get(gstin[:2], "")
    return ""


def _state_code_from_name(state_name):
    if not state_name:
        return ""
    return STATE_NAME_TO_CODE.get(state_name.strip(), "")


def _date_str(dt):
    if not dt:
        return ""
    if hasattr(dt, "strftime"):
        return dt.strftime("%d-%m-%Y")
    return str(dt)


def _fmt(value, decimals=2):
    try:
        return f"{float(value):.{decimals}f}"
    except (TypeError, ValueError):
        return "0.00"


def monthly_range(year, month):
    import calendar
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day)


def fy_range(fy_start_year):
    return date(fy_start_year, 4, 1), date(fy_start_year + 1, 3, 31)


def _fetch_invoices(org_id, start, end, include_cancelled=False):
    q = (
        Invoice.query
        .filter(
            Invoice.organization_id == org_id,
            Invoice.issue_date >= datetime.combine(start, datetime.min.time()),
            Invoice.issue_date <= datetime.combine(end, datetime.max.time()),
        )
        .options(
            joinedload(Invoice.items),
            joinedload(Invoice.customer).joinedload(Customer.lead),
        )
    )
    if not include_cancelled:
        q = q.filter(Invoice.status != InvoiceStatus.CANCELLED)
    return q.order_by(Invoice.issue_date.asc(), Invoice.invoice_number.asc()).all()


def _customer_gstin(inv):
    cust = inv.customer
    if not cust:
        return ""
    if cust.lead and cust.lead.gst_number:
        return cust.lead.gst_number.strip().upper()
    if hasattr(cust, "gst_number") and cust.gst_number:
        return cust.gst_number.strip().upper()
    return ""


def _customer_trade_name(inv):
    cust = inv.customer
    if not cust:
        return ""
    if cust.lead and cust.lead.trade_name:
        return cust.lead.trade_name
    return cust.company or cust.name or ""


def _customer_state(inv):
    cust = inv.customer
    if not cust:
        return ""
    if cust.lead and cust.lead.state:
        return cust.lead.state.strip()
    return ""


def _pos_code(inv):
    gstin = _customer_gstin(inv)
    if gstin and len(gstin) >= 2:
        return gstin[:2]
    return _state_code_from_name(_customer_state(inv))


def _is_b2b(inv):
    return bool(_customer_gstin(inv))


def _is_exempt(inv):
    if not inv.items:
        return False
    return all((it.gst_percentage or 0) == 0 for it in inv.items)


def _is_b2cl(inv):
    if _is_b2b(inv):
        return False
    return inv.is_igst and (inv.total_amount or 0) > 250000


def _is_b2cs(inv):
    if _is_b2b(inv):
        return False
    if _is_exempt(inv):
        return False
    return not _is_b2cl(inv)


def build_b2b(invoices):
    rows = []
    for inv in invoices:
        if not _is_b2b(inv):
            continue
        gstin = _customer_gstin(inv)
        receiver = _customer_trade_name(inv)
        pos = _pos_code(inv)
        for it in inv.items:
            rows.append({
                "GSTIN/UIN of Recipient": gstin,
                "Receiver Name": receiver,
                "Invoice Number": inv.invoice_number,
                "Invoice date": _date_str(inv.issue_date),
                "Invoice Value": _fmt(inv.total_amount),
                "Place Of Supply": pos,
                "Reverse Charge": "N",
                "Applicable % of Tax Rate": "",
                "Invoice Type": "Regular",
                "E-Commerce GSTIN": "",
                "Rate": _fmt(it.gst_percentage or 0),
                "Taxable Value": _fmt(it.amount or 0),
                "Cess Amount": "0",
            })
    return rows


def build_b2cl(invoices):
    rows = []
    for inv in invoices:
        if not _is_b2cl(inv):
            continue
        pos = _pos_code(inv)
        for it in inv.items:
            rows.append({
                "Invoice Number": inv.invoice_number,
                "Invoice date": _date_str(inv.issue_date),
                "Invoice Value": _fmt(inv.total_amount),
                "Place Of Supply": pos,
                "Applicable % of Tax Rate": "",
                "Rate": _fmt(it.gst_percentage or 0),
                "Taxable Value": _fmt(it.amount or 0),
                "Cess Amount": "0",
                "E-Commerce GSTIN": "",
            })
    return rows


def build_b2cs(invoices):
    buckets = {}
    for inv in invoices:
        if not _is_b2cs(inv):
            continue
        pos = _pos_code(inv)
        supply_type = "INTR" if inv.is_igst else "INRS"
        for it in inv.items:
            rate = float(it.gst_percentage or 0)
            key = (supply_type, pos, rate)
            if key not in buckets:
                buckets[key] = {
                    "Type": supply_type,
                    "Place Of Supply": pos,
                    "Applicable % of Tax Rate": "",
                    "Rate": _fmt(rate),
                    "Taxable Value": 0.0,
                    "Cess Amount": 0.0,
                    "E-Commerce GSTIN": "",
                }
            buckets[key]["Taxable Value"] += float(it.amount or 0)
    result = []
    for row in buckets.values():
        row["Taxable Value"] = _fmt(row["Taxable Value"])
        row["Cess Amount"] = _fmt(row["Cess Amount"])
        result.append(row)
    return result


def build_cdnr(invoices):
    return []


def build_cdnur(invoices):
    return []


def build_exp(invoices):
    return []


def build_exemp(invoices):
    exempt = 0.0
    nil_rated = 0.0
    non_gst = 0.0
    for inv in invoices:
        if not _is_exempt(inv):
            continue
        for it in inv.items:
            exempt += float(it.amount or 0)
    if exempt == 0 and nil_rated == 0 and non_gst == 0:
        return []
    return [{
        "Description": "Inter-State supplies to registered persons",
        "Nil Rated Supplies": _fmt(nil_rated),
        "Exempted (other than nil rated/non-GST supply)": _fmt(exempt),
        "Non-GST Supplies": _fmt(non_gst),
    }]


def _hsn_rows(invoices, b2b_only):
    buckets = {}
    for inv in invoices:
        is_b2b = _is_b2b(inv)
        if b2b_only and not is_b2b:
            continue
        if not b2b_only and is_b2b:
            continue
        for it in inv.items:
            hsn = (it.hsn_sac or "").strip()
            rate = float(it.gst_percentage or 0)
            key = (hsn, it.unit or "", rate)
            if key not in buckets:
                buckets[key] = {
                    "HSN": hsn,
                    "Description": it.item_name or it.description or "",
                    "UQC": _uqc(it.unit),
                    "Total Quantity": 0.0,
                    "Total Value": 0.0,
                    "Taxable Value": 0.0,
                    "Integrated Tax Amount": 0.0,
                    "Central Tax Amount": 0.0,
                    "State/UT Tax Amount": 0.0,
                    "Cess Amount": 0.0,
                }
            qty = float((it.chargeable_quantity or it.quantity) or 0)
            amt = float(it.amount or 0)
            gst_amt = float(it.gst_amount or 0)
            buckets[key]["Total Quantity"] += qty
            buckets[key]["Total Value"] += amt + gst_amt
            buckets[key]["Taxable Value"] += amt
            if inv.is_igst:
                buckets[key]["Integrated Tax Amount"] += gst_amt
            else:
                buckets[key]["Central Tax Amount"] += gst_amt / 2
                buckets[key]["State/UT Tax Amount"] += gst_amt / 2
    result = []
    for row in buckets.values():
        for k in ["Total Quantity", "Total Value", "Taxable Value", "Integrated Tax Amount",
                  "Central Tax Amount", "State/UT Tax Amount", "Cess Amount"]:
            row[k] = _fmt(row[k])
        result.append(row)
    return result


def build_hsn_b2b(invoices):
    return _hsn_rows(invoices, b2b_only=True)


def build_hsn_b2c(invoices):
    return _hsn_rows(invoices, b2b_only=False)


def build_docs(invoices, cancelled_invoices):
    rows = []
    if invoices:
        nums = sorted([inv.invoice_number for inv in invoices])
        rows.append({
            "Nature of Document": "Invoices for outward supply",
            "Sr. No. From": nums[0],
            "Sr. No. To": nums[-1],
            "Total Number": str(len(invoices)),
            "Cancelled": str(len(cancelled_invoices)),
        })
    return rows


def build_cancelled(all_invoices):
    rows = []
    for inv in all_invoices:
        if inv.status != InvoiceStatus.CANCELLED:
            continue
        rows.append({
            "Invoice No.": inv.invoice_number,
            "Invoice Date": _date_str(inv.issue_date),
            "Customer Name": inv.customer.name if inv.customer else "",
            "GSTIN": _customer_gstin(inv),
            "Invoice Value": _fmt(inv.total_amount),
            "Status": "Cancelled",
        })
    return rows


def build_gst_report(org_id, start, end):
    valid_invoices = _fetch_invoices(org_id, start, end, include_cancelled=False)
    all_invoices = _fetch_invoices(org_id, start, end, include_cancelled=True)
    cancelled_invoices = [inv for inv in all_invoices if inv.status == InvoiceStatus.CANCELLED]

    total_taxable = sum(float(inv.amount or 0) for inv in valid_invoices)
    total_gst = sum(float(inv.gst_amount or 0) for inv in valid_invoices)
    total_invoice_value = sum(float(inv.total_amount or 0) for inv in valid_invoices)

    return {
        "summary": {
            "total_invoices": len(valid_invoices),
            "total_taxable": _fmt(total_taxable),
            "total_gst": _fmt(total_gst),
            "total_invoice_value": _fmt(total_invoice_value),
            "cancelled_count": len(cancelled_invoices),
            "start": _date_str(start),
            "end": _date_str(end),
        },
        "b2b": build_b2b(valid_invoices),
        "b2cl": build_b2cl(valid_invoices),
        "b2cs": build_b2cs(valid_invoices),
        "cdnr": build_cdnr(valid_invoices),
        "cdnur": build_cdnur(valid_invoices),
        "exp": build_exp(valid_invoices),
        "exemp": build_exemp(valid_invoices),
        "hsn_b2b": build_hsn_b2b(valid_invoices),
        "hsn_b2c": build_hsn_b2c(valid_invoices),
        "docs": build_docs(valid_invoices, cancelled_invoices),
        "cancelled": build_cancelled(all_invoices),
    }


def build_excel(report, settings, period_label):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
    ALT_FILL = PatternFill("solid", fgColor="EEF2F7")
    BORDER = Border(
        left=Side(style="thin", color="D0D7E3"),
        right=Side(style="thin", color="D0D7E3"),
        top=Side(style="thin", color="D0D7E3"),
        bottom=Side(style="thin", color="D0D7E3"),
    )

    def _make_sheet(name, rows):
        ws = wb.create_sheet(title=name)
        if not rows:
            ws.append([f"No data available for {name}"])
            ws["A1"].font = Font(italic=True, color="888888")
            return ws
        headers = list(rows[0].keys())
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
        ws.row_dimensions[1].height = 30
        for row_idx, row in enumerate(rows, 2):
            fill = ALT_FILL if row_idx % 2 == 0 else None
            for col_idx, key in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if fill:
                    cell.fill = fill
        for col_idx, header in enumerate(headers, 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(header))
            for row in rows:
                val = str(row.get(header, "") or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)
        ws.freeze_panes = "A2"
        return ws

    _make_sheet("b2b,sez,de", report["b2b"])
    _make_sheet("b2cl", report["b2cl"])
    _make_sheet("b2cs", report["b2cs"])
    _make_sheet("cdnr", report["cdnr"])
    _make_sheet("cdnur", report["cdnur"])
    _make_sheet("exp", report["exp"])
    _make_sheet("exemp", report["exemp"])
    _make_sheet("hsn(b2b)", report["hsn_b2b"])
    _make_sheet("hsn(b2c)", report["hsn_b2c"])
    _make_sheet("docs", report["docs"])
    _make_sheet("Cancelled", report["cancelled"])

    # Summary sheet at position 0
    summary_ws = wb.create_sheet(title="Summary", index=0)
    summary_ws["A1"] = "GSTR-1 Report"
    summary_ws["A1"].font = Font(bold=True, size=16, color="1E3A5F")
    summary_ws["A2"] = f"Organisation: {settings.company_name or 'N/A'}"
    summary_ws["A2"].font = Font(bold=True, size=11, color="1E3A5F")
    summary_ws["A3"] = f"GSTIN: {settings.company_gstin or 'N/A'}"
    summary_ws["A3"].font = Font(bold=True, size=11, color="1E3A5F")
    summary_ws["A4"] = f"Period: {period_label}"
    summary_ws["A5"] = f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    rows_data = [
        ["", ""],
        ["Total Invoices (Valid)", report["summary"]["total_invoices"]],
        ["Total Taxable Value", report["summary"]["total_taxable"]],
        ["Total GST Amount", report["summary"]["total_gst"]],
        ["Total Invoice Value", report["summary"]["total_invoice_value"]],
        ["Cancelled Invoices", report["summary"]["cancelled_count"]],
    ]
    for r in rows_data:
        summary_ws.append(r)
    summary_ws.column_dimensions["A"].width = 30
    summary_ws.column_dimensions["B"].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
